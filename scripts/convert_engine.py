import sys
import os
import argparse

def convert_pdf_to_word(input_path, output_path):
    from pdf2docx import Converter
    cv = Converter(input_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()

def convert_pdf_to_excel(input_path, output_path):
    import pdfplumber
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet

    with pdfplumber.open(input_path) as pdf:
        if len(pdf.pages) == 0:
            raise ValueError("PDF has no pages")

        for i, page in enumerate(pdf.pages):
            ws = wb.create_sheet(title=f"Page {i + 1}" if len(pdf.pages) > 1 else "Sheet1")
            
            # Extract tables using both lattice (explicit lines) and stream strategies
            tables = page.extract_tables({
                "vertical_strategy": "lines",
                "horizontal_strategy": "lines",
                "snap_tolerance": 3,
                "join_tolerance": 3,
            })
            
            # Fallback to text alignment stream strategy if no bordered tables found
            if not tables or len(tables) == 0 or len(tables[0]) == 0:
                tables = page.extract_tables({
                    "vertical_strategy": "text",
                    "horizontal_strategy": "text",
                    "snap_tolerance": 4,
                    "join_tolerance": 4,
                })

            if tables and len(tables) > 0:
                current_row = 1
                for table in tables:
                    for r_idx, row in enumerate(table):
                        for c_idx, cell in enumerate(row):
                            val = cell.strip() if cell else ""
                            cell_obj = ws.cell(row=current_row, column=c_idx + 1)
                            
                            # Parse numeric/currency values
                            if val.startswith("$"):
                                try:
                                    clean_val = float(val.replace("$", "").replace(",", ""))
                                    cell_obj.value = clean_val
                                    cell_obj.number_format = "$#,##0.00"
                                except ValueError:
                                    cell_obj.value = val
                            elif val.replace(".", "", 1).replace("-", "", 1).isdigit():
                                try:
                                    cell_obj.value = int(val) if val.isdigit() else float(val)
                                except ValueError:
                                    cell_obj.value = val
                            else:
                                cell_obj.value = val

                            # Header row styling
                            if r_idx == 0:
                                cell_obj.font = Font(bold=True, color="1E3A8A")
                                cell_obj.fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")

                        current_row += 1
                    current_row += 1 # Empty row between tables
            else:
                # Fallback: extract positioned words as lines
                text = page.extract_text()
                if text:
                    for r_idx, line in enumerate(text.split("\n")):
                        ws.cell(row=r_idx + 1, column=1, value=line.strip())

            # Auto-fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)

def main():
    parser = argparse.ArgumentParser(description="Professional Document Converter Engine")
    parser.add_argument("mode", choices=["pdf2docx", "pdf2excel"])
    parser.add_argument("input", help="Input file path")
    parser.add_argument("output", help="Output file path")

    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"Error: Input file '{args.input}' not found", file=sys.stderr)
        sys.exit(1)

    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)

    try:
        if args.mode == "pdf2docx":
            convert_pdf_to_word(args.input, args.output)
        elif args.mode == "pdf2excel":
            convert_pdf_to_excel(args.input, args.output)
        print("SUCCESS")
    except Exception as e:
        print(f"Error: {str(e)}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
